"""Lee el libro de revisión del catálogo de entidades y escribe las decisiones que usa entity_catalogue.py.

Uso: python3 scripts/apply_entity_review.py Revision_Catalogo_Entidades.xlsx && pnpm run data
Salida: data-src/mission/entity_review.csv y entity_nature_review.csv (se acumulan: una decisión nueva sustituye a la anterior).
"""
import csv, os, sys
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'data-src', 'mission')
wb = load_workbook(sys.argv[1], data_only=True)

def sheet_rows(name):
    w = wb[name]; head = [c.value for c in w[1]]
    for row in w.iter_rows(min_row=2, values_only=True):
        yield dict(zip(head, row))

def merge(path, key, new):
    old = {}
    if os.path.exists(path):
        for r in csv.DictReader(open(path, encoding='utf-8')): old[key(r)] = r
    for r in new: old[key(r)] = r
    if not old: return 0
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=list(next(iter(old.values())).keys())); w.writeheader(); w.writerows(old.values())
    return len(new)

MAP1 = {'Misma entidad': 'confirmar', 'Distinta': 'rechazar'}
pairs, nat = [], []
for name in wb.sheetnames:
    w = wb[name]
    head = [c.value for c in w[1]]
    # parejas: hojas con pair/left_id/right_id y una columna de decisión (libro v1: "Decisión"; v2: "Decisión final")
    if 'pair' in head:
        dcol = 'Decisión final' if 'Decisión final' in head else 'Decisión'
        for r in sheet_rows(name):
            d = MAP1.get((r.get(dcol) or '').strip())
            if d and r.get('pair'):
                pairs.append({'pair': r['pair'], 'left_id': str(r['left_id']), 'right_id': str(r['right_id']), 'decision': d,
                              'comment': r.get('Comentario') or '', 'reviewer': r.get('Revisado por') or ''})
    # tipos: hojas con id y "Tipo final" (v2) o "Naturaleza corregida" (v1)
    elif 'id' in head and ('Tipo final' in head or 'Naturaleza corregida' in head):
        tcol = 'Tipo final' if 'Tipo final' in head else 'Naturaleza corregida'
        for r in sheet_rows(name):
            v = (r.get(tcol) or '').strip()
            if v: nat.append({'id': str(r['id']), 'nature': v, 'level': '', 'comment': r.get('Comentario') or '', 'reviewer': r.get('Revisado por') or ''})

k1 = lambda r: (r['pair'], r['left_id'], r['right_id'])
print('cruces:', merge(os.path.join(SRC, 'entity_review.csv'), k1, pairs), '| naturaleza:', merge(os.path.join(SRC, 'entity_nature_review.csv'), lambda r: r['id'], nat))
