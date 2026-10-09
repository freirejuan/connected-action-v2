"""Libro de revisión: autoridades del anexo 5 que aparecen con nombres distintos (pareja A-A), con la propuesta de un agente.

Uso: python3 scripts/entity_aa_workbook.py <carpeta agent_review> <salida.xlsx>   (lee aa.json y aa_b*_out.json)
Se aplica con scripts/apply_entity_review.py (hojas con columna pair = "A-A").
"""
import glob, json, os, sys
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

D, OUT = sys.argv[1], sys.argv[2]
items = json.load(open(os.path.join(D, 'aa.json'), encoding='utf-8'))
res = {}
for f in sorted(glob.glob(os.path.join(D, 'aa_b*_out.json'))):
    for o in json.load(open(f, encoding='utf-8')): res[o['k']] = o
assert len(res) == len(items), (len(res), len(items))

wb = Workbook()
ws = wb.active; ws.title = 'Cómo revisar'
for r in [
    ['Revisión: la misma autoridad con nombres distintos en el anexo 5'],
    ['inViable · 9-oct-2026 · Connected Action v2 · Territories'],
    [],
    ['Qué es', 'Cada proyecto escribe a su manera las autoridades del anexo 5 (Midtjylland, Central Denmark Region, Central-Jutland). Se han comparado las entradas del mismo país con el mismo código NUTS o con nombres casi iguales, y un agente ha propuesto si son la misma autoridad.'],
    ['Qué hacer', 'Hoja 1: decidir las dudas (la propuesta, si la hay, ya está puesta en "Decisión final"). Hoja 2: repaso rápido; cambiar solo lo que no compartas.'],
    ['Señal', '"Proyectos en común" no vacío casi siempre significa autoridades distintas: un proyecto no repite la misma autoridad.'],
]: ws.append(r)
ws['A1'].font = Font(bold=True, size=14); ws['A2'].font = Font(italic=True, color='666666')
ws.column_dimensions['A'].width = 16; ws.column_dimensions['B'].width = 120
for row in ws.iter_rows(min_row=4):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    if row[0].value: row[0].font = Font(bold=True)
lists = wb.create_sheet('Listas'); [lists.cell(i, 1, v) for i, v in enumerate(['Misma entidad', 'Distinta', 'Dudoso'], 1)]; lists.sheet_state = 'hidden'

COLS = ['País', 'Entrada 1', 'Código(s) 1', 'Proyectos 1', 'Entrada 2', 'Código(s) 2', 'Proyectos 2', 'Proyectos en común',
        'Propuesta del agente', 'Confianza', 'Motivo', 'Decisión final', 'Comentario', 'pair', 'left_id', 'right_id']
W = {'Entrada 1': 30, 'Entrada 2': 30, 'Proyectos 1': 26, 'Proyectos 2': 26, 'Proyectos en común': 18, 'Motivo': 46, 'Decisión final': 16, 'Comentario': 24, 'Propuesta del agente': 15}
def sheet(name, rows):
    w = wb.create_sheet(name); w.append(COLS)
    for i, c in enumerate(w[1], 1):
        c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='EEEADF'); c.alignment = Alignment(wrap_text=True, vertical='top')
        w.column_dimensions[get_column_letter(i)].width = W.get(COLS[i - 1], 12)
    w.freeze_panes = 'C2'
    for it in rows:
        o = res[it['k']]
        w.append([it['country'], it['name1'], ', '.join(it['codes1']), ', '.join(it['projects1']), it['name2'], ', '.join(it['codes2']),
                  ', '.join(it['projects2']), ', '.join(it['shared_projects']), o['decision'], o['confidence'], o['reason'],
                  o['decision'] if o['decision'] != 'Dudoso' else None, None, 'A-A', it['left_id'], it['right_id']])
        for c in w[w.max_row]: c.alignment = Alignment(wrap_text=True, vertical='top')
        if o['decision'] == 'Misma entidad': w.cell(w.max_row, 9).font = Font(bold=True)
        if o['confidence'] != 'alta': w.cell(w.max_row, 10).fill = PatternFill('solid', fgColor='FFF4CC')
    dv = DataValidation(type='list', formula1='=Listas!$A$1:$A$3', allow_blank=True, showDropDown=False)
    w.add_data_validation(dv); dv.add(f'L2:L{max(w.max_row, 2)}')
    for col in 'NOP': w.column_dimensions[col].hidden = True
    w.auto_filter.ref = w.dimensions
doubt = [it for it in items if res[it['k']]['decision'] == 'Dudoso' or res[it['k']]['confidence'] != 'alta']
ok = [it for it in items if it not in doubt]
key = lambda it: (res[it['k']]['decision'] != 'Misma entidad', it['country'], it['name1'])
sheet('1 Dudas', sorted(doubt, key=key)); sheet('2 Resueltas', sorted(ok, key=key))
wb.move_sheet('Listas', offset=len(wb.sheetnames)); wb.active = 0
wb.save(OUT)
c = Counter(res[it['k']]['decision'] for it in items)
print(OUT, '| parejas', len(items), dict(c), '| dudas', len(doubt), 'resueltas', len(ok))
