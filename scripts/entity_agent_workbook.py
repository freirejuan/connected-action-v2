"""Libro de revisión v2 del catálogo de entidades: decisiones propuestas por agentes con búsqueda web.

Entradas: <carpeta>/pairs.json, nature.json (lo que se envió a los agentes) y *_out.json (sus respuestas),
          data-src/mission/entity_candidates.csv (cruces automáticos).
Salida:   el .xlsx indicado. Se aplica con scripts/apply_entity_review.py.

Uso: python3 scripts/entity_agent_workbook.py <carpeta agent_review> <salida.xlsx>
"""
import csv, glob, json, os, re, sys
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
D, OUT = sys.argv[1], sys.argv[2]
def outs(pat):
    m = {}
    for f in sorted(glob.glob(os.path.join(D, pat))):
        for o in json.load(open(f, encoding='utf-8')): m[o['k']] = o
    return m
pin = json.load(open(os.path.join(D, 'pairs.json'), encoding='utf-8'))
nin = json.load(open(os.path.join(D, 'nature.json'), encoding='utf-8'))
P, N = outs('*pairs*_out.json'), outs('*nature*_out.json')
assert len(P) == len(pin) and len(N) == len(nin), (len(P), len(pin), len(N), len(nin))

NATURES = ['Autoridad local', 'Autoridad regional', 'Autoridad nacional', 'Otro organismo público', 'Universidad',
           'Centro de investigación', 'Empresa', 'ONG / fundación / asociación', 'Otra']
PAIR_DEC = ['Misma entidad', 'Distinta', 'Dudoso']
# criterio pendiente: asociaciones empresariales o sectoriales clasificadas como ONG (un agente las pasó a "Otra")
SECTOR = re.compile(r'empresari|sectorial|industri|asegurador|patronal|cl[uú]ster|c[aá]mara de comercio|de empresas|utilities|servicios p[uú]blicos municipales', re.I)

wb = Workbook()
lists = wb.active; lists.title = 'Listas'
for i, v in enumerate(PAIR_DEC, 1): lists.cell(i, 1, v)
for i, v in enumerate(NATURES, 1): lists.cell(i, 2, v)
for i, v in enumerate(['Distinta', 'Dudoso'], 1): lists.cell(i, 3, v)
lists.sheet_state = 'hidden'

def dropdown(ws, col, rng):
    dv = DataValidation(type='list', formula1=rng, allow_blank=True, showDropDown=False, showErrorMessage=True,
                        errorTitle='Valor no válido', error='Elige un valor de la lista')
    ws.add_data_validation(dv); dv.add(f'{col}2:{col}{max(ws.max_row, 2)}')

def header(ws, cols, widths, freeze='C2'):
    ws.append(cols)
    for i, c in enumerate(ws[1], 1):
        c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='EEEADF'); c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.column_dimensions[get_column_letter(i)].width = widths.get(cols[i - 1], 14)
    ws.freeze_panes = freeze

YELLOW = PatternFill('solid', fgColor='FFF4CC'); ORANGE = PatternFill('solid', fgColor='FCE4D6')
def style_row(ws, conf_col, conf):
    for c in ws[ws.max_row]: c.alignment = Alignment(wrap_text=True, vertical='top')
    if conf == 'media': ws.cell(ws.max_row, conf_col).fill = YELLOW
    if conf == 'baja': ws.cell(ws.max_row, conf_col).fill = ORANGE

# ---- cómo revisar ----
ws = wb.create_sheet('Cómo revisar', 0)
txt = [
    ['Revisión del catálogo de entidades · v2 (con propuestas de agentes)'],
    ['Inviable · 8-oct-2026 · Connected Action v2 · Territories'],
    [],
    ['Qué ha cambiado', 'Cada pareja dudosa y cada tipo de entidad lo ha revisado un agente con el nombre, el país, los códigos NUTS y, cuando hacía falta, búsqueda web. Su propuesta ya está en la columna de decisión final: solo hay que cambiarla si no estás de acuerdo.'],
    ['Por dónde empezar', 'Hojas 1 y 3 (dudas): decidir cada fila. Después, hojas 2 y 4 (resueltas): repaso rápido, empezando por las filas con confianza media (en amarillo). La hoja 5 son los cruces automáticos ya aplicados: marcar "Distinta" solo si alguno está mal.'],
    ['Desplegables', 'Las columnas de decisión tienen una lista desplegable en Excel y Google Sheets. Numbers (Mac) no muestra estas listas: en ese caso, escribir el valor exacto (aparece en la columna de la propuesta).'],
    ['Criterio pendiente', 'Asociaciones empresariales o sectoriales (aseguradoras, empresas de servicios públicos, clústeres): un agente las clasificó como "Otra" y los demás como "ONG / fundación / asociación". Están en las dudas de la hoja 3 con la nota "criterio": decidir una vez y aplicar igual a todas.'],
    ['Cómo se aplica', 'Al devolver el libro: python3 scripts/apply_entity_review.py <libro> && pnpm run data. Las decisiones se guardan en data-src/mission/entity_review.csv y entity_nature_review.csv.'],
]
for r in txt: ws.append(r)
ws['A1'].font = Font(bold=True, size=14); ws['A2'].font = Font(italic=True, color='666666')
ws.column_dimensions['A'].width = 22; ws.column_dimensions['B'].width = 120
for row in ws.iter_rows(min_row=4):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    if row[0].value: row[0].font = Font(bold=True)

# ---- parejas ----
PCOLS = ['Cruce', 'País', 'Entidad 1', 'Fuente 1', 'Código(s) 1', 'Entidad 2', 'Fuente 2', 'Código / sede 2', 'Propuesta del agente',
         'Confianza', 'Motivo', 'Fuente consultada', 'Decisión final', 'Comentario', 'pair', 'left_id', 'right_id']
PW = {'Entidad 1': 30, 'Entidad 2': 40, 'Fuente 1': 18, 'Fuente 2': 24, 'Motivo': 50, 'Fuente consultada': 28, 'Decisión final': 18,
      'Propuesta del agente': 16, 'Comentario': 26}
def pair_sheet(name, items, prefill):
    w = wb.create_sheet(name); header(w, PCOLS, PW)
    for p in items:
        o = P[p['k']]
        w.append([p['cross'], p['country'], p['entity1'], p['source1'], p['codes1'], p['entity2'], p['source2'], p['codes2'],
                  o['decision'], o['confidence'], o['reason'], o.get('evidence_url') or None,
                  (o['decision'] if prefill and o['decision'] != 'Dudoso' else None), None, p['pair'], p['left_id'], p['right_id']])
        style_row(w, 10, o['confidence'])
    dropdown(w, 'M', "=Listas!$A$1:$A$3")
    for col in 'OPQ': w.column_dimensions[col].hidden = True
    w.auto_filter.ref = w.dimensions
    return w
order = {'baja': 0, 'media': 1, 'alta': 2}
p_doubt = [p for p in pin if P[p['k']]['decision'] == 'Dudoso' or P[p['k']]['confidence'] != 'alta']
p_ok = [p for p in pin if p not in p_doubt]
p_doubt.sort(key=lambda p: (order[P[p['k']]['confidence']], p['cross'], p['country']))
p_ok.sort(key=lambda p: (P[p['k']]['decision'], p['cross'], p['country']))
pair_sheet('1 Parejas · dudas', p_doubt, prefill=True)
pair_sheet('2 Parejas · resueltas', p_ok, prefill=True)

# ---- tipos ----
NCOLS = ['País', 'Nombre (CORDIS)', 'Nombre corto', 'Ciudad', 'Tipo CORDIS', 'Propuesta automática', 'Propuesta del agente',
         'Confianza', 'Motivo', 'Fuente consultada', 'Tipo final', 'Comentario', 'id']
NW = {'Nombre (CORDIS)': 48, 'Nombre corto': 14, 'Propuesta automática': 22, 'Propuesta del agente': 26, 'Motivo': 46,
      'Fuente consultada': 26, 'Tipo final': 26, 'Comentario': 24}
def is_criterion(n):
    o = N[n['k']]
    return o['nature'] in ('ONG / fundación / asociación', 'Otra') and bool(SECTOR.search(o['reason'] + ' ' + (n['name'] or '')))
def nature_sheet(name, items):
    w = wb.create_sheet(name); header(w, NCOLS, NW)
    for n in items:
        o = N[n['k']]
        reason = ('[criterio] ' if is_criterion(n) else '') + o['reason']
        w.append([n['country'], n['name'], n['short'], n['city'], n['cordis_type'], n['proposal'], o['nature'], o['confidence'],
                  reason, o.get('evidence_url') or None, o['nature'], None, n['id']])
        style_row(w, 8, o['confidence'])
        if o['nature'] != n['proposal']: w.cell(w.max_row, 7).font = Font(bold=True)
    dropdown(w, 'K', "=Listas!$B$1:$B$9")
    w.column_dimensions['M'].hidden = True
    w.auto_filter.ref = w.dimensions
n_doubt = [n for n in nin if N[n['k']]['confidence'] == 'baja' or is_criterion(n)]
n_ok = [n for n in nin if n not in n_doubt]
n_doubt.sort(key=lambda n: (not is_criterion(n), n['country'], n['name']))
n_ok.sort(key=lambda n: (order[N[n['k']]['confidence']], N[n['k']]['nature'], n['country'], n['name']))
nature_sheet('3 Tipos · dudas', n_doubt)
nature_sheet('4 Tipos · resueltos', n_ok)

# ---- cruces automáticos (sin cambios) ----
A = {a['id']: a for a in json.load(open(os.path.join(ROOT, 'public/data/territories.json')))}
E = {e['id']: e for e in json.load(open(os.path.join(ROOT, 'public/data/eea_signatories_ids.json')))}
C = {e['id']: e for e in json.load(open(os.path.join(ROOT, 'public/data/entities.json')))}
nm = lambda k, i: A[i]['name'] if k == 'A' else E[i]['name'] if k == 'E' else C[i]['legal_name']
cd = lambda k, i: ', '.join(A[i]['codes']) if k == 'A' else ', '.join(E[i]['codes']) if k == 'E' else (C[i].get('related_nuts_code_nuts_code') or '')
PAIRN = {'A-C': 'Anexo 5 ↔ CORDIS', 'E-A': 'EEA ↔ Anexo 5', 'E-C': 'EEA ↔ CORDIS'}
w = wb.create_sheet('5 Cruces automáticos')
header(w, ['Cruce', 'Entidad 1', 'Código(s) 1', 'Entidad 2', 'Código / sede 2', 'Decisión final', 'Comentario', 'pair', 'left_id', 'right_id'],
       {'Entidad 1': 36, 'Entidad 2': 46, 'Decisión final': 16, 'Comentario': 26}, freeze='B2')
for r in csv.DictReader(open(os.path.join(ROOT, 'data-src/mission/entity_candidates.csv'), encoding='utf-8')):
    if r['accepted'] != 'True': continue
    l, rr = r['pair'].split('-')
    w.append([PAIRN[r['pair']], nm(l, r['left_id']), cd(l, r['left_id']), nm(rr, r['right_id']), cd(rr, r['right_id']), None, None, r['pair'], r['left_id'], r['right_id']])
dropdown(w, 'F', "=Listas!$C$1:$C$2")
for col in 'HIJ': w.column_dimensions[col].hidden = True
w.auto_filter.ref = w.dimensions

# ---- resumen ----
s = wb.create_sheet('Resumen', 1)
pc = Counter(P[p['k']]['decision'] for p in pin)
rows = [['Qué', 'Número', 'Nota'],
        ['Parejas revisadas por agente', len(pin), f"{pc['Misma entidad']} misma entidad · {pc['Distinta']} distinta · {pc['Dudoso']} dudosa"],
        ['  · dudas (hoja 1)', len(p_doubt), 'dudosas o con confianza media o baja'],
        ['  · resueltas (hoja 2)', len(p_ok), 'confianza alta'],
        ['Tipos revisados por agente', len(nin), f"{sum(1 for n in nin if N[n['k']]['nature'] != n['proposal'])} cambian respecto a la propuesta automática"],
        ['  · dudas (hoja 3)', len(n_doubt), f"{sum(1 for n in n_doubt if is_criterion(n))} por el criterio de asociaciones sectoriales, el resto con confianza baja"],
        ['  · resueltos (hoja 4)', len(n_ok), f"{sum(1 for n in n_ok if N[n['k']]['confidence'] == 'media')} con confianza media (amarillo)"],
        ['Cruces automáticos (hoja 5)', w.max_row - 1, 'ya aplicados en el Lab'],
        [], ['Tipo propuesto por el agente', 'Entidades', '']] + [[k, v, ''] for k, v in Counter(N[n['k']]['nature'] for n in nin).most_common()]
for r in rows: s.append(r)
for c in s[1]: c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='EEEADF')
s.cell(10, 1).font = Font(bold=True)
s.column_dimensions['A'].width = 34; s.column_dimensions['B'].width = 12; s.column_dimensions['C'].width = 80

wb.move_sheet('Listas', offset=len(wb.sheetnames))
wb.active = 0
wb.save(OUT)
print(OUT, '| parejas dudas', len(p_doubt), 'resueltas', len(p_ok), '| tipos dudas', len(n_doubt), 'resueltos', len(n_ok))
