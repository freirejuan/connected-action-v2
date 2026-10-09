"""Libro de revisión manual del catálogo de entidades (Territories v2).

Entrada: data-src/mission/entity_candidates.csv y entity_nature_auto.csv (de entity_catalogue.py).
Salida: el .xlsx indicado como argumento. Las decisiones se leen con apply_entity_review.py.
"""
import csv, json, os, sys
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PUB = os.path.join(ROOT, 'public', 'data'); SRC = os.path.join(ROOT, 'data-src', 'mission')
load = lambda n: json.load(open(os.path.join(PUB, n + '.json')))
A = {a['id']: a for a in load('territories')}
E = {e['id']: e for e in load('eea_signatories_ids')}
C = {e['id']: e for e in load('entities')}
types = {t['id']: t['name'] for t in load('entity_types')}
actors = load('actors')
cand = list(csv.DictReader(open(os.path.join(SRC, 'entity_candidates.csv'), encoding='utf-8')))
nat = list(csv.DictReader(open(os.path.join(SRC, 'entity_nature_auto.csv'), encoding='utf-8')))

PAIR = {'A-C': 'Anexo 5 ↔ CORDIS', 'E-A': 'EEA ↔ Anexo 5', 'E-C': 'EEA ↔ CORDIS'}
SRCN = {'A': 'anexo 5', 'E': 'EEA', 'C': 'CORDIS'}
NATURES = ['Autoridad local', 'Autoridad regional', 'Autoridad nacional', 'Otro organismo público', 'Universidad',
           'Centro de investigación', 'Empresa', 'ONG / fundación / asociación', 'Otra']

def ent(k, i):
    if k == 'A': a = A[i]; return a['name'], ', '.join(a['codes']), 'autoridad (anexo 5)' + (' · firmante' if a['is_signatory'] else '')
    if k == 'E': e = E[i]; return e['name'], ', '.join(e['codes']), f"firmante EEA ({e['level']})"
    c = C[i]; nm = c['legal_name'] + (f" ({c['short_name']})" if c.get('short_name') and c['short_name'] != c['legal_name'] else '')
    return nm, c.get('related_nuts_code_nuts_code') or '', 'socio CORDIS · ' + (types.get(c['organization_activity_type_id']) or 'Other')

def note(r):
    s, g = float(r['score_sort']), r['geo'] == 'True'
    if s == 100 and not g: return 'Mismo nombre, códigos distintos: si es la misma, uno de los dos códigos puede estar mal'
    if not g: return 'Nombre parecido, códigos distintos'
    if s < 70: return 'Un nombre contiene al otro: puede ser una entidad distinta del mismo territorio'
    return 'Nombre parecido, mismo territorio'

wb = Workbook()
ws = wb.active; ws.title = 'Cómo revisar'
txt = [
    ['Revisión del catálogo de entidades (Connected Action v2 · Territories)'],
    ['inViable · 8-oct-2026 · propuestas automáticas de scripts/entity_catalogue.py'],
    [],
    ['Para qué', 'El Lab une en una sola entidad las autoridades del anexo 5, los firmantes de la EEA y los socios de CORDIS, y clasifica cada entidad por su naturaleza. Sin esta unión no se puede saber quién firma y además participa, ni distinguir autoridades, universidades, empresas y ONG.'],
    ['Qué está hecho', 'Los cruces de confianza alta (mismo país, mismo territorio y mismo nombre tras normalizar, p. ej. "Prato" = "COMUNE DI PRATO") ya se aplican en el Lab. La naturaleza de universidades, centros de investigación y empresas sale del tipo de CORDIS y no se revisa.'],
    ['Hoja 1 · Cruces a revisar', 'Pares dudosos (confianza media o baja). Para cada fila, elegir en "Decisión": Misma entidad / Distinta / Dudoso. Las filas del mismo color son candidatos para la misma entidad: como mucho uno puede ser "Misma entidad".'],
    ['Hoja 2 · Cruces automáticos', 'Los cruces ya aplicados, para un repaso rápido. Solo hay que marcar "Distinta" si alguno está mal.'],
    ['Hoja 3 · Naturaleza', 'Socios CORDIS de tipo "organismo público" u "otros" que no están ya unidos a una autoridad. Solo hay que rellenar "Naturaleza corregida" cuando la propuesta esté mal. Las de confianza baja van primero.'],
    ['Naturalezas', 'Autoridad local (municipio, ciudad, área metropolitana) · Autoridad regional (región, provincia, condado) · Autoridad nacional (ministerio, gobierno) · Otro organismo público (agencias, institutos, consorcios públicos) · Universidad · Centro de investigación · Empresa · ONG / fundación / asociación · Otra (clústeres, redes mixtas y lo que no encaje).'],
    ['Cómo se aplica', 'Al devolver el libro, scripts/apply_entity_review.py lee las decisiones y regenera el catálogo y las vistas del Lab.'],
]
for r in txt: ws.append(r)
ws['A1'].font = Font(bold=True, size=14); ws['A2'].font = Font(italic=True, color='666666')
ws.column_dimensions['A'].width = 28; ws.column_dimensions['B'].width = 120
for row in ws.iter_rows(min_row=4):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    if row[0].value: row[0].font = Font(bold=True)

def header(w, cols, widths):
    w.append(cols)
    for i, c in enumerate(w[1], 1):
        c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='EEEADF'); c.alignment = Alignment(wrap_text=True, vertical='top')
        w.column_dimensions[get_column_letter(i)].width = widths.get(cols[i - 1], 14)
    w.freeze_panes = 'C2'

accepted_left = {(r['pair'], r['left_id']) for r in cand if r['accepted'] == 'True'}
accepted_right = {(r['pair'], r['right_id']) for r in cand if r['accepted'] == 'True'}
COLS = ['Cruce', 'Confianza', 'País', 'Entidad 1', 'Fuente 1', 'Código(s) 1', 'Entidad 2', 'Fuente 2', 'Código / sede 2',
        'Similitud', 'Qué mirar', 'Decisión', 'Comentario', 'Revisado por', 'pair', 'left_id', 'right_id']
W = {'Entidad 1': 36, 'Entidad 2': 46, 'Fuente 1': 22, 'Fuente 2': 30, 'Qué mirar': 44, 'Decisión': 18, 'Comentario': 30, 'Cruce': 16}

# 1. a revisar
w = wb.create_sheet('1 Cruces a revisar'); header(w, COLS, W)
# en una actualización solo entra lo nuevo: lo ya revisado (entity_review.csv, entity_nature_review.csv) no se repite
import csv as _csv
_rev = os.path.join(SRC, 'entity_review.csv'); _nrev = os.path.join(SRC, 'entity_nature_review.csv')
reviewed = {(r['pair'], r['left_id'], r['right_id']) for r in _csv.DictReader(open(_rev, encoding='utf-8'))} if os.path.exists(_rev) else set()
reviewed_nat = {r['id'] for r in _csv.DictReader(open(_nrev, encoding='utf-8'))} if os.path.exists(_nrev) else set()
todo = [r for r in cand if (r['pair'], r['left_id'], r['right_id']) not in reviewed and r['accepted'] != 'True' and r['confidence'] in ('media', 'baja')
        and (r['pair'], r['left_id']) not in accepted_left and (r['pair'], r['right_id']) not in accepted_right]
order = {'media': 0, 'baja': 1}
todo.sort(key=lambda r: (r['pair'], order[r['confidence']], A.get(r['left_id'], E.get(r['left_id'], {})).get('country', ''), r['left_id'], int(r['rank'])))
shade = ['FFFFFF', 'F3F0EA']; g = -1; last = None
for r in todo:
    l, rr = r['pair'].split('-')
    n1, c1, s1 = ent(l, r['left_id']); n2, c2, s2 = ent(rr, r['right_id'])
    country = (A.get(r['left_id']) or E.get(r['left_id']))['country']
    if (r['pair'], r['left_id']) != last: g += 1; last = (r['pair'], r['left_id'])
    w.append([PAIR[r['pair']], r['confidence'], country, n1, s1, c1, n2, s2, c2, round(float(r['score_sort'])), note(r), None, None, None, r['pair'], r['left_id'], r['right_id']])
    for c in w[w.max_row]:
        c.alignment = Alignment(wrap_text=True, vertical='top'); c.fill = PatternFill('solid', fgColor=shade[g % 2])
dv = DataValidation(type='list', formula1='"Misma entidad,Distinta,Dudoso"', allow_blank=True); w.add_data_validation(dv); dv.add(f'L2:L{max(w.max_row, 2)}')
for col in 'OPQ': w.column_dimensions[col].hidden = True
w.auto_filter.ref = w.dimensions
n_todo = len(todo); n_todo_left = len({(r['pair'], r['left_id']) for r in todo})

# 2. automáticos
w = wb.create_sheet('2 Cruces automáticos'); header(w, COLS, W)
auto = [r for r in cand if r['accepted'] == 'True']
auto.sort(key=lambda r: (r['pair'], float(r['score_sort'])))
for r in auto:
    l, rr = r['pair'].split('-')
    n1, c1, s1 = ent(l, r['left_id']); n2, c2, s2 = ent(rr, r['right_id'])
    country = (A.get(r['left_id']) or E.get(r['left_id']))['country']
    w.append([PAIR[r['pair']], r['confidence'], country, n1, s1, c1, n2, s2, c2, round(float(r['score_sort'])), 'Ya aplicado: marcar "Distinta" solo si está mal', None, None, None, r['pair'], r['left_id'], r['right_id']])
    for c in w[w.max_row]: c.alignment = Alignment(wrap_text=True, vertical='top')
dv = DataValidation(type='list', formula1='"Distinta,Dudoso"', allow_blank=True); w.add_data_validation(dv); dv.add(f'L2:L{max(w.max_row, 2)}')
for col in 'OPQ': w.column_dimensions[col].hidden = True
w.auto_filter.ref = w.dimensions

# 3. naturaleza
in_authority = {cid for a in actors if a['annex_ids'] or a['eea_ids'] for cid in a['cordis_ids']}
w = wb.create_sheet('3 Naturaleza')
cols = ['Confianza', 'País', 'Nombre (CORDIS)', 'Nombre corto', 'Tipo CORDIS', 'Propuesta', 'Por qué', 'Naturaleza corregida', 'Comentario', 'Revisado por', 'id']
header(w, cols, {'Nombre (CORDIS)': 60, 'Nombre corto': 18, 'Tipo CORDIS': 16, 'Propuesta': 26, 'Por qué': 28, 'Naturaleza corregida': 28, 'Comentario': 30})
rows = [n for n in nat if n['cordis_type'] in ('Public bodies', 'Other') and n['id'] not in in_authority and n['id'] not in reviewed_nat]
rows.sort(key=lambda n: ({'baja': 0, 'media': 1}.get(n['confidence'], 2), n['country'], n['name']))
TYPES_ES = {'Public bodies': 'Organismo público', 'Other': 'Otros'}
for n in rows:
    w.append([n['confidence'], n['country'], n['name'], n['short_name'], TYPES_ES[n['cordis_type']], n['nature'], n['rule'], None, None, None, n['id']])
    for c in w[w.max_row]: c.alignment = Alignment(wrap_text=True, vertical='top')
    if n['confidence'] == 'baja': w.cell(w.max_row, 1).fill = PatternFill('solid', fgColor='FCE4D6')
dv = DataValidation(type='list', formula1='"' + ','.join(NATURES) + '"', allow_blank=True); w.add_data_validation(dv); dv.add(f'H2:H{max(w.max_row, 2)}')
w.column_dimensions['K'].hidden = True
w.auto_filter.ref = w.dimensions

# resumen al principio
w = wb.create_sheet('Resumen', 1)
from collections import Counter
nat_c = Counter(a['nature'] for a in actors)
summary = [
    ['Qué', 'Número', 'Nota'],
    ['Entidades unificadas', len(actors), 'autoridades del anexo 5 + firmantes EEA + socios CORDIS, una vez unidos'],
    ['Cruces automáticos ya aplicados', len(auto), ' · '.join(f"{PAIR[k]}: {v}" for k, v in sorted(Counter(r['pair'] for r in auto).items()))],
    ['Pares a revisar (hoja 1)', n_todo, f'{n_todo_left} entidades con uno o más candidatos dudosos'],
    ['Naturalezas a revisar (hoja 3)', len(rows), f"{sum(1 for n in rows if n['confidence'] == 'baja')} de confianza baja"],
    ['Firman y participan', sum(1 for a in actors if a['signatory'] and (a['annex_ids'] or a['cordis_ids'])), 'con los cruces automáticos; cambiará con la revisión'],
    ['Solo firman', sum(1 for a in actors if a['signatory'] and not a['annex_ids'] and not a['cordis_ids']), 'firmantes EEA sin pareja en el anexo 5 ni en CORDIS'],
    [],
    ['Naturaleza (propuesta actual)', 'Entidades', ''],
] + [[k, v, ''] for k, v in nat_c.most_common()]
for r in summary: w.append(r)
for c in w[1]: c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='EEEADF')
w.cell(10, 1).font = Font(bold=True)
w.column_dimensions['A'].width = 34; w.column_dimensions['B'].width = 12; w.column_dimensions['C'].width = 90

out = sys.argv[1] if len(sys.argv) > 1 else 'Revision_Catalogo_Entidades.xlsx'
wb.save(out)
print(out, '| a revisar', n_todo, 'pares /', n_todo_left, 'entidades | automáticos', len(auto), '| naturaleza', len(rows))
